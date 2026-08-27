"""Generate Product module Test Plan + Test Cases Excel workbook."""
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    import subprocess
    import sys

    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parents[1] / "Product_Module_Test_Plan_and_TestCases.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SECTION_FILL = PatternFill("solid", fgColor="D6EAF8")
PASS_FILL = PatternFill("solid", fgColor="E8F5E9")
THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)


def style_header(ws, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(1, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def add_rows(ws, rows):
    for r_idx, row in enumerate(rows, 2):
        for c_idx, val in enumerate(row, 1):
            cell = ws.cell(r_idx, c_idx, val)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = THIN


# ---------------------------------------------------------------------------
# Sheet 1: Test Plan
# ---------------------------------------------------------------------------
plan_headers = [
    "Section",
    "Item",
    "Details",
]

plan_rows = [
    ["Document Info", "Module", "Inventory → Products (New Products)"],
    ["Document Info", "URL", "https://quickvee.com/merchants/inventory/new-products"],
    ["Document Info", "Add Single", "/merchants/inventory/new-products/add?type=single"],
    ["Document Info", "Add Variants", "/merchants/inventory/new-products/add?type=variants"],
    ["Document Info", "Edit", "/merchants/inventory/new-products/edit/{productId}"],
    ["Document Info", "Branch", "Product_Feature"],
    ["Document Info", "Environment", "Production (quickvee.com) — also validate on QA"],
    ["Document Info", "Store used for discovery", "Gang Smoker (gang)"],
    ["Document Info", "Priority", "P0 — Core commerce module"],
    [
        "Scope In",
        "In Scope",
        "Product list, search/filter/sort, create single, create variants, edit, delete, "
        "row actions (view details, sales history, instant PO, stocktake, online ordering toggles), "
        "photos, pricing, inventory, taxes, SKU/custom codes, selling channels, multi-store copy, "
        "related products, vendor assignment on edit, validations, empty/error states, permissions",
    ],
    [
        "Scope Out",
        "Out of Scope (separate modules)",
        "Categories/Brands/Tags/Attributes/Vendors CRUD pages themselves; Bulk Updates page; "
        "Import Data; Inventory Export; Transfers; Lottery; Purchase Order full flow; Stocktake full app — "
        "covered only as entry points from Products",
    ],
    [
        "Objectives",
        "Goal",
        "Ensure merchants can create, find, update, sell, and control products across register + online "
        "channels with correct inventory, pricing, taxes, and variants without data loss or invalid saves",
    ],
    [
        "Test Strategy",
        "Levels",
        "UI smoke + functional (CRUD) + field validation + negative + regression + multi-store + "
        "permission + API response sanity where automation exists",
    ],
    [
        "Test Strategy",
        "Types",
        "Positive, Negative, Boundary, UI/UX, Usability, Integration (Brand/Category/Tag/Vendor/Attribute/Tax), "
        "Compatibility (desktop breakpoints), Accessibility basics",
    ],
    [
        "Entry Criteria",
        "Ready when",
        "Product_Feature branch deployed; test store has categories/brands/attributes/taxes/vendors seeded; "
        "login works on quickvee.com/login",
    ],
    [
        "Exit Criteria",
        "Done when",
        "All P0/P1 cases Pass or waived; no open Blocker/Critical bugs; smoke suite green",
    ],
    [
        "Risks",
        "High",
        "Cannot switch Single ↔ Variants after create (must delete & re-add). Wrong type = major merchant pain.",
    ],
    [
        "Risks",
        "High",
        "Custom codes/UPCs must be unique store-wide; collisions break scanning at POS",
    ],
    [
        "Risks",
        "Medium",
        "Variant combination explosion (up to 3 attributes) — performance & pricing validation",
    ],
    [
        "Risks",
        "Medium",
        "Multi-store copy may create inconsistent inventory/channel settings across outlets",
    ],
    [
        "Dependencies",
        "Master data",
        "Categories*, Brands, Tags, Attributes (saved), Vendors, Taxes, Linked stores",
    ],
    [
        "Roles",
        "Permission",
        "Merchant admin vs employee with limited inventory rights; verify Access Denied on Add product when denied",
    ],
    [
        "List UI Inventory",
        "Header",
        "Products title; 'N loaded · scroll for more'; Search; Add product",
    ],
    [
        "List UI Inventory",
        "Filters",
        "Categories | Brands | Tags | Vendors | Channels | No photo | Status",
    ],
    [
        "List UI Inventory",
        "Toolbar",
        "Online ordering (bulk enable/disable delivery/pickup) | Sort",
    ],
    [
        "List UI Inventory",
        "Columns",
        "Product (name, UPC/SKU, categories/brand) | Availability (qty/Low/Inactive + Delivery/Pickup icons) | Price (or by variant)",
    ],
    [
        "List UI Inventory",
        "Row actions",
        "Edit product | View details | Sales history | Instant Purchase Order | Stocktake | Delete item | "
        "Online Ordering Delivery/Pickup toggles",
    ],
    [
        "Create — Type Step",
        "Step 1 of 2",
        "Single product (One UPC, One price, Single inventory) OR Product with variants "
        "(Per-variant UPC/price/inventory). Warning: irreversible after create. Continue disabled until type selected.",
    ],
    [
        "Create/Edit — Single fields",
        "Product Information",
        "Name* | Brand (one) | Categories* | Photos (first = cover) | Description (rich text B/I/U, max 2000, Online only)",
    ],
    [
        "Create/Edit — Single fields",
        "Pricing & Inventory",
        "UPC + Generate | Cost | Price* | Compare-at | Margin% (computed) | Profit$ (computed) | "
        "Available to sell | Reorder point | Reorder quantity",
    ],
    [
        "Create/Edit — Single fields",
        "Copy to stores (create only)",
        "Select all / per-store checkboxes for linked stores (e.g. Chain Smoker, Fun, Rocker)",
    ],
    [
        "Create/Edit — Single fields",
        "Selling Channels",
        "POS always on | Delivery | Pickup (defaults enabled on new)",
    ],
    [
        "Create/Edit — Single fields",
        "Product Options",
        "Track quantity | Continue selling when out of stock | Check ID (Tobacco) | Active | Food Stampable (SNAP/EBT)",
    ],
    [
        "Create/Edit — Single fields",
        "SKU Codes",
        "Custom Code + Add another (up to 15 additional scan codes; unique store-wide; shared uniqueness with UPCs)",
    ],
    [
        "Create/Edit — Single fields",
        "Vendor Information",
        "Assign after create on edit: vendors + cost + preferred vendor",
    ],
    [
        "Create/Edit — Single fields",
        "Tax Information",
        "Default/extra taxes (name + rate); Add another tax; empty = store default",
    ],
    [
        "Create/Edit — Single fields",
        "Related Products",
        "Suggested alongside this product online",
    ],
    [
        "Create/Edit — Variants extras",
        "Attributes",
        "Up to 3 attributes from saved Attributes; values via Enter; combinations auto-generated; every variant needs price",
    ],
    [
        "Create/Edit — Variants extras",
        "Per variant",
        "Expand variant for inventory, options, codes (UPC/price/qty per variant)",
    ],
    [
        "Sort options",
        "Sort",
        "Newest first | Oldest first | Name A→Z | Name Z→A | Category A→Z | Category Z→A",
    ],
    [
        "Channel filter",
        "Channels",
        "Online Pickup | Online Delivery | Select all / Clear | Cancel | Apply",
    ],
    [
        "Status filter",
        "Status",
        "All statuses | Pending | Approved | Rejected",
    ],
    [
        "Online ordering bulk",
        "Actions",
        "Enable All (Delivery+Pickup) | Enable Delivery for All | Enable Pickup for All | "
        "Disable Delivery for All | Disable Pickup for All | Disable All",
    ],
    [
        "Search",
        "Placeholder",
        "Search by product name, variant name, UPC, SKU, or custom code",
    ],
]

# ---------------------------------------------------------------------------
# Sheet 2: Test Cases
# ---------------------------------------------------------------------------
tc_headers = [
    "TC ID",
    "Module / Area",
    "Feature",
    "Priority",
    "Type",
    "Preconditions",
    "Test Steps",
    "Test Data",
    "Expected Result",
    "Status",
    "Actual Result",
    "Bug ID",
    "Notes",
]


def tc(
    tid,
    area,
    feature,
    priority,
    ttype,
    pre,
    steps,
    data,
    expected,
    notes="",
):
    return [
        tid,
        area,
        feature,
        priority,
        ttype,
        pre,
        steps,
        data,
        expected,
        "Not Executed",
        "",
        "",
        notes,
    ]


cases = []

# --- Navigation / Access ---
cases += [
    tc(
        "PRD-001",
        "Navigation",
        "Open Products from Inventory menu",
        "P0",
        "Smoke",
        "Logged in merchant",
        "1. Open sidebar Inventory\n2. Click Products",
        "URL: /merchants/inventory/new-products",
        "Products list loads; search, filters, Add product visible; infinite scroll hint shown",
    ),
    tc(
        "PRD-002",
        "Navigation",
        "Direct URL access",
        "P1",
        "Functional",
        "Logged in",
        "1. Navigate directly to products URL",
        "https://quickvee.com/merchants/inventory/new-products",
        "Same list page renders without error",
    ),
    tc(
        "PRD-003",
        "Permissions",
        "Add product Access Denied",
        "P0",
        "Negative",
        "Employee without product create permission",
        "1. Login as restricted user\n2. Open Products\n3. Hover/click Add product",
        "Restricted role",
        "Add product blocked / Access Denied indicator; cannot open create flow",
        "UI shows Access Denied wrapper on Add product",
    ),
    tc(
        "PRD-004",
        "Permissions",
        "View-only can open list but not delete",
        "P1",
        "Negative",
        "View-only inventory role",
        "1. Open row actions\n2. Attempt Delete",
        "View-only user",
        "Delete hidden or blocked with permission message",
    ),
]

# --- List UI ---
cases += [
    tc(
        "PRD-010",
        "Product List",
        "UI elements display",
        "P0",
        "UI",
        "Products exist",
        "1. Open Products page\n2. Verify header, search, filters, sort, columns, rows",
        "N/A",
        "Visible: Products title, loaded count, Search, Add product, Filters "
        "(Categories/Brands/Tags/Vendors/Channels/No photo/Status), Online ordering, Sort, "
        "columns Product/Availability/Price, rows with Show actions",
    ),
    tc(
        "PRD-011",
        "Product List",
        "Empty catalog state",
        "P1",
        "UI",
        "Store with 0 products (or filtered to none)",
        "1. Apply filter that matches nothing OR use empty store",
        "N/A",
        "Friendly empty state; Add product still available; no console crash",
    ),
    tc(
        "PRD-012",
        "Product List",
        "Infinite scroll load more",
        "P0",
        "Functional",
        "Catalog has > page size products",
        "1. Note 'N loaded · scroll for more'\n2. Scroll to bottom\n3. Wait for more rows",
        "Store with many products",
        "Additional products load; count increases; no duplicates; no stuck loading",
    ),
    tc(
        "PRD-013",
        "Product List",
        "Row displays single product fields",
        "P0",
        "UI",
        "Single product exists",
        "1. Locate a single product row",
        "e.g. gang quickadd",
        "Shows name, UPC/code, category/brand, qty or Low · N, Delivery/Pickup icons, price",
    ),
    tc(
        "PRD-014",
        "Product List",
        "Row displays variant product fields",
        "P0",
        "UI",
        "Variant product exists",
        "1. Locate a variant product row",
        "e.g. product with 'by variant'",
        "Shows name, categories, variant count, Delivery/Pickup, price with 'by variant'",
    ),
    tc(
        "PRD-015",
        "Product List",
        "Inactive product indicator",
        "P1",
        "UI",
        "Inactive product exists",
        "1. Find inactive product",
        "e.g. Vivek Product",
        "Availability shows Inactive; price may show —; Delivery/Pickup off",
    ),
    tc(
        "PRD-016",
        "Product List",
        "Low stock indicator",
        "P1",
        "UI",
        "Product at/below reorder point",
        "1. Find low stock product",
        "Low · N",
        "Availability shows Low · quantity",
    ),
]

# --- Search ---
cases += [
    tc(
        "PRD-020",
        "Search",
        "Search by product name",
        "P0",
        "Functional",
        "Known product name",
        "1. Type full/partial name in search\n2. Wait for results",
        "Name: gang quickadd",
        "Matching products listed; non-matches hidden",
    ),
    tc(
        "PRD-021",
        "Search",
        "Search by UPC",
        "P0",
        "Functional",
        "Product with UPC",
        "1. Enter UPC in search",
        "Valid UPC",
        "Correct product returned",
    ),
    tc(
        "PRD-022",
        "Search",
        "Search by SKU / custom code",
        "P0",
        "Functional",
        "Product with custom code",
        "1. Enter custom code",
        "Custom code value",
        "Correct product returned",
    ),
    tc(
        "PRD-023",
        "Search",
        "Search by variant name",
        "P1",
        "Functional",
        "Variant product with distinct variant names",
        "1. Search variant name",
        "Variant value name",
        "Parent variant product appears",
    ),
    tc(
        "PRD-024",
        "Search",
        "No results",
        "P1",
        "Negative",
        "Any catalog",
        "1. Search nonsense string",
        "zzzz_no_match_999",
        "Empty/no-match state; no error toast crash",
    ),
    tc(
        "PRD-025",
        "Search",
        "Clear search restores list",
        "P1",
        "Functional",
        "Search applied",
        "1. Clear search box",
        "N/A",
        "Full list (or previous filters) restored",
    ),
    tc(
        "PRD-026",
        "Search",
        "Case insensitive search",
        "P2",
        "Functional",
        "Mixed-case product name",
        "1. Search with different casing",
        "GANG QUICKADD / gang quickadd",
        "Same results",
    ),
]

# --- Filters ---
cases += [
    tc(
        "PRD-030",
        "Filters",
        "Filter by Categories",
        "P0",
        "Functional",
        "Products in multiple categories",
        "1. Open Categories filter\n2. Select one category\n3. Apply",
        "Category with known products",
        "Only products in that category shown; filter chip/state active",
    ),
    tc(
        "PRD-031",
        "Filters",
        "Filter by Brands",
        "P0",
        "Functional",
        "Products with brands",
        "1. Open Brands\n2. Select brand\n3. Apply",
        "Known brand",
        "Only matching brand products shown",
    ),
    tc(
        "PRD-032",
        "Filters",
        "Filter by Tags",
        "P1",
        "Functional",
        "Tagged products exist",
        "1. Open Tags\n2. Select tag\n3. Apply",
        "Known tag",
        "Only tagged products shown",
    ),
    tc(
        "PRD-033",
        "Filters",
        "Filter by Vendors",
        "P1",
        "Functional",
        "Vendor assigned to products",
        "1. Open Vendors\n2. Select vendor\n3. Apply",
        "Known vendor",
        "Only products supplied by vendor shown",
    ),
    tc(
        "PRD-034",
        "Filters",
        "Filter Channels — Online Pickup",
        "P0",
        "Functional",
        "Mix of pickup on/off",
        "1. Channels → Online Pickup → Apply",
        "N/A",
        "Only pickup-enabled products listed",
    ),
    tc(
        "PRD-035",
        "Filters",
        "Filter Channels — Online Delivery",
        "P0",
        "Functional",
        "Mix of delivery on/off",
        "1. Channels → Online Delivery → Apply",
        "N/A",
        "Only delivery-enabled products listed",
    ),
    tc(
        "PRD-036",
        "Filters",
        "Channels Select all / Clear / Cancel",
        "P1",
        "Functional",
        "On Channels popover",
        "1. Select all\n2. Clear\n3. Cancel without apply",
        "N/A",
        "Select all checks both; Clear resets; Cancel leaves previous filter unchanged",
    ),
    tc(
        "PRD-037",
        "Filters",
        "No photo filter",
        "P1",
        "Functional",
        "Some products without photos",
        "1. Click No photo",
        "N/A",
        "Only products without photos listed; toggle off restores",
    ),
    tc(
        "PRD-038",
        "Filters",
        "Status — All / Pending / Approved / Rejected",
        "P1",
        "Functional",
        "Products in different approval statuses if applicable",
        "1. Open Status\n2. Select each option",
        "All statuses, Pending, Approved, Rejected",
        "List filters correctly per status; All shows everything",
    ),
    tc(
        "PRD-039",
        "Filters",
        "Combined filters + search",
        "P1",
        "Functional",
        "Rich catalog",
        "1. Apply Category + Brand\n2. Add search term",
        "Valid combo",
        "AND logic applied; results satisfy all criteria",
    ),
    tc(
        "PRD-040",
        "Filters",
        "Clear all filters",
        "P1",
        "Functional",
        "Multiple filters active",
        "1. Clear each filter / reset",
        "N/A",
        "Full unfiltered list restored",
    ),
]

# --- Sort ---
cases += [
    tc(
        "PRD-050",
        "Sort",
        "Newest first (default)",
        "P0",
        "Functional",
        "Multiple products",
        "1. Select Sort: Newest first",
        "N/A",
        "Most recently created appear first",
    ),
    tc(
        "PRD-051",
        "Sort",
        "Oldest first",
        "P1",
        "Functional",
        "Multiple products",
        "1. Select Oldest first",
        "N/A",
        "Oldest products first",
    ),
    tc(
        "PRD-052",
        "Sort",
        "Name A → Z",
        "P1",
        "Functional",
        "Multiple products",
        "1. Select Name (A → Z)",
        "N/A",
        "Alphabetical ascending by name",
    ),
    tc(
        "PRD-053",
        "Sort",
        "Name Z → A",
        "P1",
        "Functional",
        "Multiple products",
        "1. Select Name (Z → A)",
        "N/A",
        "Alphabetical descending",
    ),
    tc(
        "PRD-054",
        "Sort",
        "Category A → Z / Z → A",
        "P2",
        "Functional",
        "Products with categories",
        "1. Select Category sorts",
        "N/A",
        "Ordered by category name accordingly",
    ),
]

# --- Online ordering bulk ---
cases += [
    tc(
        "PRD-060",
        "Online Ordering Bulk",
        "Enable All (Delivery + Pickup)",
        "P0",
        "Functional",
        "Products with channels off",
        "1. Online ordering → Enable All",
        "N/A",
        "All products show Delivery on + Pickup on (or confirmation then apply)",
    ),
    tc(
        "PRD-061",
        "Online Ordering Bulk",
        "Enable Delivery for All",
        "P1",
        "Functional",
        "Some delivery off",
        "1. Enable Delivery for All",
        "N/A",
        "Delivery enabled catalog-wide; Pickup unchanged",
    ),
    tc(
        "PRD-062",
        "Online Ordering Bulk",
        "Enable Pickup for All",
        "P1",
        "Functional",
        "Some pickup off",
        "1. Enable Pickup for All",
        "N/A",
        "Pickup enabled catalog-wide; Delivery unchanged",
    ),
    tc(
        "PRD-063",
        "Online Ordering Bulk",
        "Disable Delivery / Pickup / All",
        "P0",
        "Functional",
        "Channels enabled",
        "1. Disable Delivery for All\n2. Disable Pickup for All\n3. Disable All",
        "N/A",
        "Corresponding channel icons turn off; POS still available",
    ),
]

# --- Create type selection ---
cases += [
    tc(
        "PRD-070",
        "Create — Type",
        "Open Add product dialog",
        "P0",
        "Smoke",
        "Create permission",
        "1. Click Add product",
        "N/A",
        "Dialog: Step 1 of 2 Choose product type; Single vs Variants cards; Cancel; Continue disabled",
    ),
    tc(
        "PRD-071",
        "Create — Type",
        "Continue disabled until type chosen",
        "P0",
        "Negative",
        "Dialog open",
        "1. Do not select type\n2. Observe Continue",
        "N/A",
        "Continue remains disabled",
    ),
    tc(
        "PRD-072",
        "Create — Type",
        "Select Single product → Continue",
        "P0",
        "Functional",
        "Dialog open",
        "1. Select Single product\n2. Continue",
        "N/A",
        "Navigates to /add?type=single; form shows Single product badge",
    ),
    tc(
        "PRD-073",
        "Create — Type",
        "Select Product with variants → Continue",
        "P0",
        "Functional",
        "Dialog open",
        "1. Select Product with variants\n2. Continue",
        "N/A",
        "Navigates to /add?type=variants; attribute builder visible",
    ),
    tc(
        "PRD-074",
        "Create — Type",
        "Cancel closes dialog",
        "P1",
        "Functional",
        "Dialog open",
        "1. Click Cancel or Close",
        "N/A",
        "Returns to list; no product created",
    ),
    tc(
        "PRD-075",
        "Create — Type",
        "Irreversible type warning shown",
        "P1",
        "UI",
        "Dialog open",
        "1. Read warning copy",
        "N/A",
        "Warning: once created, switching single/variants not possible — must delete and re-add",
    ),
]

# --- Single create validations & happy path ---
cases += [
    tc(
        "PRD-080",
        "Create Single",
        "Required: Name empty",
        "P0",
        "Negative",
        "On add single form",
        "1. Leave Name empty\n2. Fill Price & Category\n3. Save",
        "Name blank",
        "Validation error on Name*; product not saved",
    ),
    tc(
        "PRD-081",
        "Create Single",
        "Required: Categories empty",
        "P0",
        "Negative",
        "On add single form",
        "1. Enter Name + Price\n2. Leave Categories empty\n3. Save",
        "Categories blank",
        "Validation error on Categories*; not saved",
    ),
    tc(
        "PRD-082",
        "Create Single",
        "Required: Price empty / invalid",
        "P0",
        "Negative",
        "On add single form",
        "1. Leave Price empty or enter negative/letters\n2. Save",
        "Price blank / -1 / abc",
        "Validation prevents save; clear error message",
    ),
    tc(
        "PRD-083",
        "Create Single",
        "Happy path — minimum required fields",
        "P0",
        "Functional",
        "Category exists",
        "1. Name, Categories, Price\n2. Save product",
        "Name=Auto Test Single; Price=9.99; Category=any",
        "Success; redirects/list shows product; Active; POS on; Delivery+Pickup default on",
    ),
    tc(
        "PRD-084",
        "Create Single",
        "Happy path — all fields",
        "P0",
        "Functional",
        "Brand, category, tax, photo ready",
        "1. Fill Name, Brand, Categories, Photos, Description\n"
        "2. UPC Generate, Cost, Price, Compare-at, Qty, Reorder point/qty\n"
        "3. Options toggles, Custom codes, Tax, Related products\n4. Save",
        "Complete dataset",
        "Product saved with all values persisted on edit reopen",
    ),
    tc(
        "PRD-085",
        "Create Single",
        "UPC Generate button",
        "P1",
        "Functional",
        "On form",
        "1. Click Generate beside UPC",
        "N/A",
        "Valid unique UPC populated",
    ),
    tc(
        "PRD-086",
        "Create Single",
        "Duplicate UPC rejected",
        "P0",
        "Negative",
        "Existing product UPC known",
        "1. Enter existing UPC\n2. Save",
        "Duplicate UPC",
        "Error: UPC already exists / not unique; not saved",
    ),
    tc(
        "PRD-087",
        "Create Single",
        "Margin & Profit auto-compute",
        "P1",
        "Functional",
        "On form",
        "1. Enter Cost=10, Price=20\n2. Observe Margin & Profit",
        "Cost 10 / Price 20",
        "Margin 50% (or formula used by app); Profit $10; read-only computed",
    ),
    tc(
        "PRD-088",
        "Create Single",
        "Compare-at price display logic",
        "P2",
        "Functional",
        "On form",
        "1. Set Compare-at higher than Price\n2. Save\n3. Check online/register presentation if applicable",
        "Price 10, Compare-at 15",
        "Compare-at stored; strikethrough/compare shown where designed",
    ),
    tc(
        "PRD-089",
        "Create Single",
        "Description max 2000 chars",
        "P1",
        "Boundary",
        "On form",
        "1. Enter 2000 chars — counter 2000/2000\n2. Try 2001",
        "2000 / 2001 chars",
        "2000 allowed; beyond blocked or truncated with feedback",
    ),
    tc(
        "PRD-090",
        "Create Single",
        "Description rich text B/I/U",
        "P2",
        "Functional",
        "On form",
        "1. Apply Bold/Italic/Underline\n2. Save & reopen",
        "Sample formatted text",
        "Formatting persists for online description",
    ),
    tc(
        "PRD-091",
        "Create Single",
        "Photos — add, cover = first, remove",
        "P0",
        "Functional",
        "Image files ready",
        "1. Add multiple photos\n2. Verify first is cover\n3. Reorder/remove if supported\n4. Save",
        "JPG/PNG images",
        "Photos saved; first used as cover on list/online/register",
    ),
    tc(
        "PRD-092",
        "Create Single",
        "Invalid photo type/size",
        "P1",
        "Negative",
        "On form",
        "1. Upload unsupported/huge file",
        "exe / 20MB+",
        "Clear error; form remains usable",
    ),
    tc(
        "PRD-093",
        "Create Single",
        "Brand — one brand only",
        "P1",
        "Functional",
        "Multiple brands exist",
        "1. Select a brand\n2. Try selecting second if UI allows",
        "Brand A",
        "Only one brand associated; helper text 'One brand per product'",
    ),
    tc(
        "PRD-094",
        "Create Single",
        "Multiple categories",
        "P1",
        "Functional",
        "Multiple categories",
        "1. Assign 2+ categories\n2. Save",
        "Cat1, Cat2",
        "All categories saved; list shows them",
    ),
    tc(
        "PRD-095",
        "Create Single",
        "Custom codes — add up to 15",
        "P1",
        "Boundary",
        "On form",
        "1. Add codes via + Add another code up to 15\n2. Try 16th",
        "15 unique codes",
        "15 accepted; 16th blocked",
    ),
    tc(
        "PRD-096",
        "Create Single",
        "Custom code uniqueness vs UPC",
        "P0",
        "Negative",
        "Existing UPC/code in store",
        "1. Set custom code = existing UPC of another product\n2. Save",
        "Colliding code",
        "Rejected — unique across store shared with UPCs",
    ),
    tc(
        "PRD-097",
        "Create Single",
        "Track quantity off",
        "P1",
        "Functional",
        "On form",
        "1. Turn off Track quantity\n2. Save",
        "N/A",
        "Qty fields disabled/hidden as designed; product sells without stock tracking",
    ),
    tc(
        "PRD-098",
        "Create Single",
        "Continue selling when out of stock",
        "P1",
        "Functional",
        "Track qty on",
        "1. Enable continue selling\n2. Set qty 0\n3. Save\n4. Attempt order/POS (if in scope)",
        "Qty 0",
        "Orders still allowed when out of stock",
    ),
    tc(
        "PRD-099",
        "Create Single",
        "Check ID (Tobacco) flag",
        "P1",
        "Functional",
        "On form",
        "1. Enable Check ID\n2. Save\n3. Verify register age-check behavior if testable",
        "N/A",
        "Flag persisted; register requires age verification for item",
    ),
    tc(
        "PRD-100",
        "Create Single",
        "Active unchecked (create inactive)",
        "P0",
        "Functional",
        "On form",
        "1. Uncheck Active\n2. Save",
        "N/A",
        "Product saved Inactive; disabled everywhere per copy",
    ),
    tc(
        "PRD-101",
        "Create Single",
        "Food Stampable SNAP/EBT",
        "P2",
        "Functional",
        "On form",
        "1. Enable Food Stampable\n2. Save",
        "N/A",
        "Flag persisted for EBT-eligible checkout paths",
    ),
    tc(
        "PRD-102",
        "Create Single",
        "Selling channels Delivery/Pickup toggles",
        "P0",
        "Functional",
        "On form",
        "1. Toggle Delivery off, Pickup on\n2. Save\n3. Verify list icons",
        "N/A",
        "POS always on; list shows Delivery off / Pickup on",
    ),
    tc(
        "PRD-103",
        "Create Single",
        "Tax — default + add another",
        "P1",
        "Functional",
        "Multiple taxes configured",
        "1. Keep DefaultTax\n2. Add another tax\n3. Save",
        "DefaultTax 10.63% + extra",
        "Taxes persisted on product",
    ),
    tc(
        "PRD-104",
        "Create Single",
        "Related products assignment",
        "P2",
        "Functional",
        "Other products exist",
        "1. Add related products\n2. Save\n3. Check online suggestions if possible",
        "2 related SKUs",
        "Related products saved",
    ),
    tc(
        "PRD-105",
        "Create Single",
        "Copy to stores — select stores",
        "P0",
        "Functional",
        "Linked multi-store merchant",
        "1. Check one linked store\n2. Save\n3. Switch store and verify product exists",
        "Copy to Chain Smoker",
        "Product created in selected stores",
    ),
    tc(
        "PRD-106",
        "Create Single",
        "Copy to stores — Select all",
        "P1",
        "Functional",
        "Multiple linked stores",
        "1. Select all\n2. Save\n3. Verify each store",
        "All linked stores",
        "Product present in all selected outlets",
    ),
    tc(
        "PRD-107",
        "Create Single",
        "Vendor info locked until after create",
        "P1",
        "UI",
        "On create form",
        "1. Scroll to Vendor Information",
        "N/A",
        "Message: vendors assignable after product is created — save first then reopen",
    ),
    tc(
        "PRD-108",
        "Create Single",
        "Back / discard without save",
        "P1",
        "Functional",
        "Dirty form",
        "1. Enter data\n2. Click Back",
        "Unsaved changes",
        "Confirm discard or navigate away without creating product",
    ),
    tc(
        "PRD-109",
        "Create Single",
        "Price = 0 allowed?",
        "P2",
        "Boundary",
        "On form",
        "1. Set Price 0\n2. Save",
        "Price 0",
        "Document actual: either allowed (free item) or validation error — capture expected from BA",
        "Confirm with product owner",
    ),
    tc(
        "PRD-110",
        "Create Single",
        "Very long product name",
        "P2",
        "Boundary",
        "On form",
        "1. Enter extremely long name\n2. Save",
        "255+ chars",
        "Max-length enforced or truncated with message; list UI does not break",
    ),
]

# --- Variant create ---
cases += [
    tc(
        "PRD-120",
        "Create Variants",
        "Attribute picker from saved attributes",
        "P0",
        "Functional",
        "Attributes exist (Size/Color/Flavor)",
        "1. Open variant create\n2. Choose attribute\n3. Add values with Enter",
        "Attribute=Size; Values=S,M,L",
        "Values accepted; combinations generated (3)",
    ),
    tc(
        "PRD-121",
        "Create Variants",
        "Up to 3 attributes",
        "P0",
        "Boundary",
        "On variant form",
        "1. Add 3 attributes with values\n2. Try 4th",
        "3 attributes",
        "Max 3 allowed; 4th blocked",
    ),
    tc(
        "PRD-122",
        "Create Variants",
        "Combination generation",
        "P0",
        "Functional",
        "2 attributes × 2 values",
        "1. Attr1: A,B\n2. Attr2: 1,2\n3. Observe variants",
        "2×2",
        "4 combinations generated; message shows count",
    ),
    tc(
        "PRD-123",
        "Create Variants",
        "Every variant needs a price",
        "P0",
        "Negative",
        "Variants generated",
        "1. Leave a variant price empty\n2. Save",
        "Missing price on one variant",
        "Validation blocks save; highlights missing prices",
    ),
    tc(
        "PRD-124",
        "Create Variants",
        "Per-variant UPC / inventory / options / codes",
        "P0",
        "Functional",
        "Variants generated",
        "1. Expand each variant\n2. Set UPC, qty, price, options\n3. Save",
        "Unique UPC per variant",
        "Each variant stores independent UPC/price/inventory",
    ),
    tc(
        "PRD-125",
        "Create Variants",
        "Happy path save variant product",
        "P0",
        "Functional",
        "Name, Categories, attributes ready",
        "1. Fill Name*, Categories*\n2. Build attributes\n3. Price all variants\n4. Save",
        "Name=Auto Variant Tee",
        "Product on list with 'by variant' and variant count; editable",
    ),
    tc(
        "PRD-126",
        "Create Variants",
        "Cannot save with 0 combinations",
        "P0",
        "Negative",
        "No attributes/values",
        "1. Fill Name/Categories only\n2. Save without attributes",
        "0 combinations",
        "Blocked with guidance to pick attribute and values",
    ),
    tc(
        "PRD-127",
        "Create Variants",
        "Remove attribute / value updates combinations",
        "P1",
        "Functional",
        "Attributes added",
        "1. Remove a value\n2. Observe combination count",
        "N/A",
        "Combinations recalculate correctly",
    ),
    tc(
        "PRD-128",
        "Create Variants",
        "Duplicate attribute value rejected",
        "P2",
        "Negative",
        "Adding values",
        "1. Add value 'Red' twice",
        "Red, Red",
        "Duplicate value prevented",
    ),
    tc(
        "PRD-129",
        "Create Variants",
        "Shared description for all variants",
        "P1",
        "Functional",
        "On form",
        "1. Set description\n2. Save\n3. Verify single description applies to product",
        "Shared desc",
        "One description for parent; copy states single & variants share description",
    ),
]

# --- Edit ---
cases += [
    tc(
        "PRD-140",
        "Edit Product",
        "Open edit via row click",
        "P0",
        "Functional",
        "Product exists",
        "1. Click product row/name",
        "N/A",
        "Opens /edit/{id}; fields prefilled",
    ),
    tc(
        "PRD-141",
        "Edit Product",
        "Open edit via row action",
        "P0",
        "Functional",
        "Product exists",
        "1. Show actions → Edit product",
        "N/A",
        "Full product form opens",
    ),
    tc(
        "PRD-142",
        "Edit Product",
        "Update name/price/qty and Save changes",
        "P0",
        "Functional",
        "On edit",
        "1. Change Name, Price, Available to sell\n2. Save changes",
        "Updated values",
        "Success toast; list reflects updates",
    ),
    tc(
        "PRD-143",
        "Edit Product",
        "Discard changes",
        "P1",
        "Functional",
        "Dirty edit form",
        "1. Change fields\n2. Discard changes",
        "N/A",
        "Changes reverted; original values remain",
    ),
    tc(
        "PRD-144",
        "Edit Product",
        "Assign vendor after create",
        "P0",
        "Functional",
        "Vendors exist; product saved",
        "1. Edit product\n2. Vendor Information → Add vendor\n3. Set cost & preferred\n4. Save",
        "Vendor + cost",
        "Vendor assigned; preferred marked; list vendor filter works",
    ),
    tc(
        "PRD-145",
        "Edit Product",
        "Cannot change Single ↔ Variant type",
        "P0",
        "Negative",
        "Existing single product",
        "1. Open edit\n2. Look for type switcher",
        "N/A",
        "No way to convert type; badge shows current type only",
    ),
    tc(
        "PRD-146",
        "Edit Product",
        "Edit variant — existing values locked; add new values",
        "P0",
        "Functional",
        "Variant product",
        "1. Edit variant product\n2. Verify existing attribute values locked\n"
        "3. Add new value\n4. Price new variants\n5. Save",
        "New value e.g. XL",
        "Copy shown: Existing attribute values locked. New variants created on save; old data retained",
    ),
    tc(
        "PRD-146A",
        "Edit Product",
        "Variant Bulk edit + Select all",
        "P1",
        "Functional",
        "Variant product with 2+ variants",
        "1. Edit variants\n2. Select all\n3. Bulk edit price/cost/active\n4. Save",
        "Bulk price update",
        "All selected variants updated consistently",
    ),
    tc(
        "PRD-146B",
        "Edit Product",
        "Per-variant Active toggle",
        "P1",
        "Functional",
        "Variant product",
        "1. Disable Active on one variant\n2. Save\n3. Verify POS/online for that variant",
        "One variant inactive",
        "Inactive variant not sellable; sibling variants unaffected",
    ),
    tc(
        "PRD-147",
        "Edit Product",
        "Toggle Active off/on",
        "P0",
        "Functional",
        "Active product",
        "1. Uncheck Active → Save\n2. Re-enable → Save",
        "N/A",
        "List shows Inactive then Active; selling blocked when inactive",
    ),
    tc(
        "PRD-148",
        "Edit Product",
        "Required field validation on edit",
        "P0",
        "Negative",
        "On edit",
        "1. Clear Name or Price or Categories\n2. Save",
        "Cleared required",
        "Validation errors; changes not saved",
    ),
]

# --- Row actions ---
cases += [
    tc(
        "PRD-160",
        "Row Actions",
        "View details quick look",
        "P1",
        "Functional",
        "Product exists",
        "1. Show actions → View details",
        "N/A",
        "Quick look panel/modal without full page navigation",
    ),
    tc(
        "PRD-161",
        "Row Actions",
        "Sales history",
        "P1",
        "Functional",
        "Product with sales",
        "1. Show actions → Sales history",
        "N/A",
        "Shows units sold & revenue for product",
    ),
    tc(
        "PRD-162",
        "Row Actions",
        "Instant Purchase Order",
        "P1",
        "Functional",
        "Product exists",
        "1. Show actions → Instant Purchase Order",
        "N/A",
        "Opens quick PO flow to add inventory for product",
    ),
    tc(
        "PRD-163",
        "Row Actions",
        "Stocktake entry",
        "P1",
        "Functional",
        "Product exists",
        "1. Show actions → Stocktake",
        "N/A",
        "Opens adjust on-hand quantities flow for product",
    ),
    tc(
        "PRD-164",
        "Row Actions",
        "Delete item — confirm",
        "P0",
        "Functional",
        "Deletable test product",
        "1. Show actions → Delete item\n2. Confirm",
        "Test product",
        "Product removed from catalog; not in search",
    ),
    tc(
        "PRD-165",
        "Row Actions",
        "Delete item — cancel",
        "P1",
        "Functional",
        "Product exists",
        "1. Delete item\n2. Cancel confirm",
        "N/A",
        "Product remains",
    ),
    tc(
        "PRD-166",
        "Row Actions",
        "Toggle Delivery from row menu",
        "P0",
        "Functional",
        "Product with Delivery on",
        "1. Show actions → Online Ordering → toggle Delivery",
        "N/A",
        "Delivery icon updates on row without full edit",
    ),
    tc(
        "PRD-167",
        "Row Actions",
        "Toggle Pickup from row menu",
        "P0",
        "Functional",
        "Product with Pickup on",
        "1. Toggle Pickup",
        "N/A",
        "Pickup icon updates immediately",
    ),
]

# --- Integration / regression ---
cases += [
    tc(
        "PRD-180",
        "Integration",
        "Product appears on POS/register",
        "P0",
        "Integration",
        "Active product with POS",
        "1. Create/activate product\n2. Open register/POS\n3. Search/scan",
        "UPC/name",
        "Product found; correct price; ID check if flagged",
    ),
    tc(
        "PRD-181",
        "Integration",
        "Online storefront visibility",
        "P0",
        "Integration",
        "Delivery/Pickup enabled + Active + photo optional",
        "1. Enable online channels\n2. Open View Online Store\n3. Find product",
        "Online store URL",
        "Product visible online with description/photos; hidden when channel off or inactive",
    ),
    tc(
        "PRD-182",
        "Integration",
        "Sale decrements Available to sell",
        "P0",
        "Integration",
        "Track quantity on",
        "1. Note qty\n2. Complete sale of 1\n3. Refresh Products",
        "Qty before N",
        "Available to sell becomes N-1; Low badge if below reorder",
    ),
    tc(
        "PRD-183",
        "Integration",
        "Brand/Category rename reflects on product",
        "P2",
        "Regression",
        "Product linked to brand/category",
        "1. Rename brand/category in master\n2. Recheck product list/edit",
        "N/A",
        "Updated names shown on product",
    ),
    tc(
        "PRD-184",
        "Integration",
        "Deleted category/brand handling",
        "P1",
        "Negative",
        "Product assigned",
        "1. Delete category used only by product (if allowed)\n2. Open product",
        "N/A",
        "Graceful handling — validation or orphan cleared; no crash",
    ),
    tc(
        "PRD-185",
        "Integration",
        "Attribute used by variants cannot break product",
        "P1",
        "Regression",
        "Variant product using attribute",
        "1. Attempt rename/delete attribute in Attributes module\n2. Reopen product",
        "N/A",
        "Product variants remain intact or delete blocked with message",
    ),
]

# --- API / data integrity (for automation suite) ---
cases += [
    tc(
        "PRD-190",
        "API / Data",
        "Create product API success response",
        "P1",
        "API",
        "Network tab / automation",
        "1. Save product\n2. Capture create API",
        "Valid payload",
        "2xx + success message; product id returned; list refresh uses new id",
    ),
    tc(
        "PRD-191",
        "API / Data",
        "Update product API",
        "P1",
        "API",
        "Existing product",
        "1. Save changes\n2. Capture update API",
        "N/A",
        "Success; GET/list shows updated fields",
    ),
    tc(
        "PRD-192",
        "API / Data",
        "Delete product API",
        "P1",
        "API",
        "Test product",
        "1. Delete\n2. Capture API",
        "N/A",
        "Success; subsequent GET 404/not in list",
    ),
    tc(
        "PRD-193",
        "API / Data",
        "UI count vs API list count",
        "P1",
        "API",
        "On list load",
        "1. Intercept product list API\n2. Compare loaded rows / total",
        "N/A",
        "UI loaded count consistent with API pagination contract",
    ),
]

# --- UX / Non-functional ---
cases += [
    tc(
        "PRD-200",
        "UX / Performance",
        "List loads within acceptable time",
        "P2",
        "Performance",
        "Normal network",
        "1. Open Products\n2. Measure time to first rows",
        "N/A",
        "First content < 3s on broadband; loading state shown",
    ),
    tc(
        "PRD-201",
        "UX / Performance",
        "Large variant matrix save",
        "P1",
        "Performance",
        "3 attrs with several values",
        "1. Generate large combo set\n2. Bulk-price\n3. Save",
        "e.g. 3×3×3=27",
        "Save completes without timeout; all variants stored",
    ),
    tc(
        "PRD-202",
        "UX",
        "Responsive layout desktop widths",
        "P2",
        "UI",
        "Desktop",
        "1. Resize 1440 / 1280 / 1024",
        "N/A",
        "Layout usable; filters accessible; no overlapping actions",
    ),
    tc(
        "PRD-203",
        "Security",
        "XSS in product name/description",
        "P0",
        "Security",
        "On create/edit",
        "1. Enter <script>alert(1)</script> in Name/Description\n2. Save\n3. View list/online",
        "XSS strings",
        "Rendered as text; no script execution",
    ),
    tc(
        "PRD-204",
        "Security",
        "Unauthorized edit via URL",
        "P1",
        "Security",
        "Restricted user",
        "1. Open /edit/{id} directly",
        "Other store product id if possible",
        "Access denied / redirect; no data leak",
    ),
]

# --- Field matrix sheet ---
field_headers = [
    "Area",
    "Field / Control",
    "Required",
    "Type",
    "Validation / Rules",
    "Single",
    "Variants",
    "Create",
    "Edit",
    "Notes",
]

fields = [
    ["Type Step", "Single product", "Yes (one of)", "Choice", "Irreversible after create", "Y", "N", "Y", "N", "Badge only on edit"],
    ["Type Step", "Product with variants", "Yes (one of)", "Choice", "Irreversible after create", "N", "Y", "Y", "N", ""],
    ["Product Info", "Name", "Yes", "Text", "Required; max length TBD", "Y", "Y", "Y", "Y", "*"],
    ["Product Info", "Brand", "No", "Single select", "One brand per product", "Y", "Y", "Y", "Y", ""],
    ["Product Info", "Categories", "Yes", "Multi select", "At least one", "Y", "Y", "Y", "Y", "*"],
    ["Product Info", "Photos", "No", "File upload", "First = cover; register + online", "Y", "Y", "Y", "Y", ""],
    ["Product Info", "Description", "No", "Rich text", "Max 2000; Online only; B/I/U", "Y", "Y", "Y", "Y", "Shared for variants"],
    ["Pricing", "UPC", "No*", "Text + Generate", "Unique store-wide with custom codes", "Y", "Per variant", "Y", "Y", "*may be required by business"],
    ["Pricing", "Cost", "No", "Currency", ">=0", "Y", "Per variant*", "Y", "Y", ""],
    ["Pricing", "Price", "Yes", "Currency", "Required; >=0 policy TBD", "Y", "Per variant Yes", "Y", "Y", "*"],
    ["Pricing", "Compare-at price", "No", "Currency", "Usually >= Price", "Y", "Per variant*", "Y", "Y", ""],
    ["Pricing", "Margin %", "N/A", "Computed", "From price & cost", "Y", "Y", "Y", "Y", "Read-only"],
    ["Pricing", "Profit $", "N/A", "Computed", "From price & cost", "Y", "Y", "Y", "Y", "Read-only"],
    ["Inventory", "Available to sell", "No", "Number", "Integer >=0 when tracking", "Y", "Per variant", "Y", "Y", ""],
    ["Inventory", "Reorder point", "No", "Number", "Triggers Low badge", "Y", "Per variant*", "Y", "Y", ""],
    ["Inventory", "Reorder quantity", "No", "Number", "Suggested reorder qty", "Y", "Per variant*", "Y", "Y", ""],
    ["Multi-store", "Copy to stores", "No", "Checkboxes", "Linked stores only", "Y", "Y", "Y", "N", "Create only"],
    ["Channels", "Point of Sale", "Yes", "Toggle", "Always on", "Y", "Y", "Y", "Y", "Cannot disable"],
    ["Channels", "Delivery", "No", "Toggle", "Default on for new", "Y", "Y", "Y", "Y", "Also row/bulk"],
    ["Channels", "Pickup", "No", "Toggle", "Default on for new", "Y", "Y", "Y", "Y", "Also row/bulk"],
    ["Options", "Track quantity", "No", "Toggle", "Keep stock counts", "Y", "Per variant*", "Y", "Y", ""],
    ["Options", "Continue selling when out of stock", "No", "Toggle", "Allow at zero", "Y", "Y", "Y", "Y", ""],
    ["Options", "Check ID", "No", "Toggle", "Tobacco age verification", "Y", "Y", "Y", "Y", ""],
    ["Options", "Active", "No", "Toggle", "Uncheck disables everywhere", "Y", "Y", "Y", "Y", "Default on"],
    ["Options", "Food Stampable", "No", "Toggle", "SNAP/EBT eligible", "Y", "Y", "Y", "Y", ""],
    ["SKU Codes", "Custom Code(s)", "No", "Text multi", "Up to 15; unique w/ UPCs", "Y", "Per variant*", "Y", "Y", ""],
    ["Vendors", "Vendor + cost + preferred", "No", "Association", "After create on edit", "Y", "Y", "N", "Y", ""],
    ["Tax", "Tax name + rate", "No", "Multi", "Empty = store default", "Y", "Y", "Y", "Y", "Add another tax"],
    ["Related", "Related Products", "No", "Multi product", "Online suggestions", "Y", "Y", "Y", "Y", ""],
    ["Variants", "Attribute (max 3)", "Yes for variants", "Select", "From saved Attributes", "N", "Y", "Y", "Y", ""],
    ["Variants", "Attribute values", "Yes for variants", "Tags/Enter", "Generate combinations", "N", "Y", "Y", "Y", ""],
    ["Variants", "Variant row expand", "—", "Panel", "Inventory/options/codes", "N", "Y", "Y", "Y", ""],
    ["List", "Search", "—", "Text", "Name/variant/UPC/SKU/custom", "Y", "Y", "—", "—", ""],
    ["List", "Filters", "—", "Multi", "Cat/Brand/Tag/Vendor/Channels/No photo/Status", "Y", "Y", "—", "—", ""],
    ["List", "Sort", "—", "Menu", "Newest/Oldest/Name/Category", "Y", "Y", "—", "—", ""],
]

# --- Traceability ---
trace_headers = ["Requirement / Feature", "TC IDs", "Priority"]
trace_rows = [
    ["Navigate to Products list", "PRD-001, PRD-002", "P0"],
    ["Permissions / Access Denied", "PRD-003, PRD-004, PRD-204", "P0"],
    ["List UI & infinite scroll", "PRD-010–PRD-016, PRD-012", "P0"],
    ["Search", "PRD-020–PRD-026", "P0"],
    ["Filters", "PRD-030–PRD-040", "P0"],
    ["Sort", "PRD-050–PRD-054", "P1"],
    ["Bulk online ordering", "PRD-060–PRD-063", "P0"],
    ["Create type selection", "PRD-070–PRD-075", "P0"],
    ["Create single product", "PRD-080–PRD-110", "P0"],
    ["Create variant product", "PRD-120–PRD-129", "P0"],
    ["Edit product", "PRD-140–PRD-148", "P0"],
    ["Row actions", "PRD-160–PRD-167", "P0"],
    ["POS / Online / Inventory integration", "PRD-180–PRD-185", "P0"],
    ["API & counts", "PRD-190–PRD-193", "P1"],
    ["Performance / Security / UX", "PRD-200–PRD-204", "P0/P2"],
]


def main():
    wb = Workbook()

    # Plan
    ws1 = wb.active
    ws1.title = "Test Plan"
    ws1.append(plan_headers)
    style_header(ws1, len(plan_headers))
    add_rows(ws1, plan_rows)
    autosize(ws1, [22, 28, 90])
    ws1.row_dimensions[1].height = 24
    ws1.freeze_panes = "A2"

    # Cases
    ws2 = wb.create_sheet("Test Cases")
    ws2.append(tc_headers)
    style_header(ws2, len(tc_headers))
    add_rows(ws2, cases)
    autosize(ws2, [12, 18, 36, 10, 12, 28, 50, 28, 45, 14, 24, 12, 24])
    ws2.row_dimensions[1].height = 30
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:M{len(cases) + 1}"

    # Fields
    ws3 = wb.create_sheet("Field Matrix")
    ws3.append(field_headers)
    style_header(ws3, len(field_headers))
    add_rows(ws3, fields)
    autosize(ws3, [16, 34, 12, 16, 40, 10, 14, 10, 10, 28])
    ws3.freeze_panes = "A2"

    # Traceability
    ws4 = wb.create_sheet("Coverage Traceability")
    ws4.append(trace_headers)
    style_header(ws4, len(trace_headers))
    add_rows(ws4, trace_rows)
    autosize(ws4, [40, 40, 12])
    ws4.freeze_panes = "A2"

    # Summary
    ws5 = wb.create_sheet("Summary")
    summary = [
        ["Metric", "Value"],
        ["Total Test Cases", len(cases)],
        ["P0 Count", sum(1 for c in cases if c[3] == "P0")],
        ["P1 Count", sum(1 for c in cases if c[3] == "P1")],
        ["P2 Count", sum(1 for c in cases if c[3] == "P2")],
        ["Module URL", "https://quickvee.com/merchants/inventory/new-products"],
        ["Login URL", "https://quickvee.com/login"],
        ["Branch", "Product_Feature"],
        ["How to use", "Filter Test Cases sheet by Priority/Module; update Status column while executing"],
    ]
    for r in summary:
        ws5.append(r)
    style_header(ws5, 2)
    for r in range(2, len(summary) + 1):
        for c in range(1, 3):
            ws5.cell(r, c).border = THIN
            ws5.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
    autosize(ws5, [24, 80])

    wb.save(OUT)
    print(f"Wrote {OUT} with {len(cases)} test cases")


if __name__ == "__main__":
    main()
